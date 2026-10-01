def read_price_row(csv_path, person_number):
    """Return the two prices from the requested line, counting from 1."""
    with open(csv_path) as price_file:
        lines = price_file.readlines()
    
    line_number = person_number - 1
    line = lines[line_number]
    starter_price, main_course_price = line.strip().split(",")

    return starter_price, main_course_price


def read_all_prices(csv_path):
    """Return all prices as a list of tuples."""
    prices = []

    with open(csv_path) as price_file:
        for line in price_file:
            starter_price, main_course_price = line.strip().split(",")
            prices.append((starter_price, main_course_price))

    return prices


def get_person_prices(all_prices, person_number):
    """Return the two prices for one person, counting from 1."""
    position_in_list = person_number - 1
    starter_price, main_course_price = all_prices[position_in_list]

    return starter_price, main_course_price


# Code for the graphical interface version.
def read_excel_lists(excel_path):
    """Read the Excel file and return the orders list and the menu list."""
    from openpyxl import load_workbook
    workbook = load_workbook(excel_path, data_only=True)

    orders_sheet = workbook["orders"] if "orders" in workbook.sheetnames else workbook["commandes"]  # Support the original workbook.
    menu_sheet = workbook["menu"]

    orders_list = []
    for row in orders_sheet.iter_rows(min_row=2, values_only=True):
        orders_list.append((row[0], row[1], row[2]))

    menu_list = []
    for row in menu_sheet.iter_rows(min_row=2, values_only=True):
        menu_list.append((row[0], row[1], row[2]))

    return orders_list, menu_list


def get_client_prices(client_number, orders_list, menu_list):
    """Return the starter price and main-course price for one client."""
    for number, starter, main_course in orders_list:
        if number == client_number:
            break

    for dish_name, category, price in menu_list:
        if dish_name == starter:
            starter_price = price
        if dish_name == main_course:
            main_course_price = price

    return starter_price, main_course_price


def calculate_price_auto(starter_price, main_course_price):
    starter_price = float(starter_price)
    main_course_price = float(main_course_price)
    total_price = starter_price + main_course_price
    return total_price


def calculate_client_total(client_number, excel_path="data/orders_menu.xlsx"):
    """Read the Excel file and return the total price for one client."""
    orders_list, menu_list = read_excel_lists(excel_path)
    starter_price, main_course_price = get_client_prices(
        client_number, orders_list, menu_list
    )
    return calculate_price_auto(starter_price, main_course_price)


def show_menu_interface(excel_path="data/orders_menu.xlsx"):
    """Display starter and main-course dropdown menus in a notebook."""
    import ipywidgets as widgets
    from IPython.display import display

    _, menu_list = read_excel_lists(excel_path)

    starter_prices = {}
    main_course_prices = {}

    for dish_name, category, price in menu_list:
        if category in ("Starter", "Entrée"):
            starter_prices[dish_name] = price
        if category in ("Main course", "Plat principal"):
            main_course_prices[dish_name] = price

    starter_dropdown = widgets.Dropdown(
        options=list(starter_prices),
        description="Starter:",
        style={"description_width": "initial"},
    )

    main_course_dropdown = widgets.Dropdown(
        options=list(main_course_prices),
        description="Main course:",
        style={"description_width": "initial"},
    )

    total_display = widgets.HTML()

    def update_total(change=None):
        starter_price = starter_prices[starter_dropdown.value]
        main_course_price = main_course_prices[main_course_dropdown.value]
        total_price = calculate_price_auto(starter_price, main_course_price)
        total_display.value = f"<b>Total price: {total_price:.2f} €</b>"

    starter_dropdown.observe(update_total, names="value")
    main_course_dropdown.observe(update_total, names="value")

    update_total()
    display(widgets.VBox([starter_dropdown, main_course_dropdown, total_display]))
